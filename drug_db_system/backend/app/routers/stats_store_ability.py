# -*- coding: utf-8 -*-
"""门店能力分析接口：/api/{db_key}/stats/store_ability（+ /export）

本期范围（业务指定）：
  - 仅大参林（dashenlin），其它库返回 400；前端按 supports_store_ability 控制菜单显隐
  - 城市固定「广州」（city_col IN ('广州','广州市')，兼容两种写法且可命中索引）
  - 门店身份按「门店名称」分组（不用门店编码）
  - 产出口径为实销盒数 SUM(数量)，不涉及金额

输出内容（一张表，默认前 100 家）：
  1. 门店总产出排行：按实销盒数降序
  2. 品类能力：每个门店的销量 Top3 品类 + 末位品类（倒数第一）
       - 门店品类数不足 3 时，多余的 Top 列留空
       - 末位列始终输出，即使与 Top3 中某一项重复（品类数 <= 3 时必然重复）
       - 占比 = 该品类盒数 / 该门店实销总数（分母与表中「实销总数」列一致）

筛选条件（均可选，留空 = 不限）：
  - stores  ：门店名称关键词，英文逗号分隔，模糊「包含」匹配（ILIKE '%词%'），多词之间 OR
  - products：品类（商品编码），英文逗号分隔，精确匹配，多词之间 IN
  - 两者留空即为「全部门店 × 全部品类」

日期范围必填（date_from / date_to 至少其一）：缺省会退化为对全量销售明细的全表扫描，
千万级数据下不可接受，因此接口直接拒绝；前端默认已带区间。首尾相同即单日查询，
用 >= 与 <= 自然命中。

GET /stats/store_ability/latest_range 返回广州「数据最新月份」的起止日期，
供前端默认填充时间范围（数据常滞后于当前自然月，不能用系统月份）。
"""
import calendar
import io
import re
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from psycopg2 import sql

from ..dependencies import validate_db_key
from ..models.schema_def import get_cfg, get_field_map, supports_store_ability
from ..services.product_map import get_product_map
from ..database import get_conn
from ..logger import get_logger

router = APIRouter(prefix="/api/{db_key}", tags=["stats-store-ability"])
logger = get_logger("app.routers.stats_store_ability")

# ---------- 本期固定口径 ----------
FIXED_CITY = "广州"
# 城市列在库里「广州」「广州市」两种写法都可能存在。
# 这里用 IN 精确枚举而不用 LIKE '广州%'：LIKE 前缀匹配在中文排序规则下无法命中
# (城市, 日期) 复合索引，千万级数据会退化成全表扫描；IN 可以正常走索引。
FIXED_CITY_VARIANTS = ("广州", "广州市")
DEFAULT_TOP_N = 100
MAX_TOP_N = 500
TOP_CATEGORY_COUNT = 3  # 取销量前 3 的品类
# 门店名称关键词个数上限：每个词生成一个 ILIKE，过多会拖慢查询
MAX_STORE_TERMS = 20
# 门店趋势按天展示的柱子上限（约 10 年），防止误选超长区间导致柱子密到无法阅读
MAX_TREND_DAYS = 3660

# ---------- Excel 样式（与 stats_store.py 保持一致） ----------
HEADER_FILL = PatternFill(start_color="3B6BD6", end_color="3B6BD6", fill_type="solid")
HEADER_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Microsoft YaHei", size=10)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)


def _parse_list(val: Optional[str]) -> List[str]:
    if not val:
        return []
    return [v.strip() for v in val.split(",") if v.strip() != ""]


def _escape_like(val: str) -> str:
    """转义 ILIKE 通配符，避免用户输入的 % / _ 被当作通配符。

    PG 的 LIKE/ILIKE 默认转义符就是反斜杠，故无需额外写 ESCAPE 子句。
    """
    return val.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _day_sequence(d_from: date, d_to: date) -> List[str]:
    """生成 [d_from, d_to] 区间内的每一天（升序），返回 YYYY-MM-DD 字符串列表。

    用于补齐没有销量的日期（qty 补 0），保证柱状图横轴连续、柱子疏密一致。
    """
    days: List[str] = []
    cur = d_from
    step = timedelta(days=1)
    while cur <= d_to:
        days.append(cur.strftime("%Y-%m-%d"))
        cur += step
    return days


def _safe_sheet_name(title: str) -> str:
    """生成合法 Excel sheet 名（去非法字符、截断 31 字符）。"""
    name = re.sub(r'[:\\/\?\*\[\]]', "_", str(title)).strip()
    return (name or "sheet1")[:31]


def _run_store_ability(
    db_key: str,
    date_from: Optional[str],
    date_to: Optional[str],
    products: Optional[str],
    top_n: int,
    stores: Optional[str] = None,
    map_names: bool = False,
) -> dict:
    """门店能力分析核心逻辑：查询 + 组装，JSON 端点与导出共用。

    stores / products 均可选，留空表示不限（全部门店 / 全部品类）。
    map_names=True 时品类显示名优先用映射表中文名，无映射回落到库内商品名称。
    """
    fm = get_field_map(db_key)
    if not fm:
        raise HTTPException(400, f"该统计功能暂不支持数据库: {db_key}")
    if not supports_store_ability(db_key):
        raise HTTPException(400, f"门店能力分析暂不支持数据库: {db_key}（本期仅大参林开放）")

    # 必须指定日期范围：否则 WHERE 不含日期条件，会退化为对全量销售明细的全表扫描，
    # 千万级数据下不可接受。前端默认已带当月区间，此校验主要堵住手动调用 /docs 的口子。
    if not date_from and not date_to:
        raise HTTPException(
            400,
            "门店能力分析必须指定日期范围（date_from 或 date_to 至少其一），"
            "以避免对全量销售明细做全表扫描。",
        )

    # 品类编码映射：仅在开关打开时读盘（文件很小，读一次足够）
    pmap: Dict[str, str] = get_product_map(db_key) if map_names else {}

    store_col = fm["store_name_col"]
    product_col = fm["product_col"]
    product_name_col = fm.get("product_name_col", "商品名称")
    qty_col = fm["qty_col"]
    city_col = fm["city_col"]
    date_col = fm["date_col"]
    table = get_cfg(db_key)["table"]

    # ---------- WHERE 条件（参数顺序必须与占位符顺序一致） ----------
    conditions: List[sql.Composable] = [
        sql.SQL("{} IN ({})").format(
            sql.Identifier(city_col),
            sql.SQL(", ").join(sql.Placeholder() * len(FIXED_CITY_VARIANTS)),
        ),
        sql.SQL("{} IS NOT NULL").format(sql.Identifier(store_col)),
        sql.SQL("btrim({}) <> ''").format(sql.Identifier(store_col)),
        sql.SQL("{} IS NOT NULL").format(sql.Identifier(date_col)),
    ]
    params: list = list(FIXED_CITY_VARIANTS)

    if date_from:
        conditions.append(sql.SQL("{} >= %s").format(sql.Identifier(date_col)))
        params.append(date_from)
    if date_to:
        conditions.append(sql.SQL("{} <= %s").format(sql.Identifier(date_col)))
        params.append(date_to)

    # 门店名称筛选：多个关键词之间 OR，模糊「包含」匹配（ILIKE '%词%'），
    # 留空则不限门店。ILIKE 本身命中不了索引，但前面的 城市 IN + 日期区间 已通过
    # (城市, 日期) 复合索引把结果集收敛到「当月的广州数据」，在此基础上做字符串
    # 匹配代价可接受；若后续门店量过大，可改用 GIN 全文索引（schema 里 门店名称
    # 已登记在 fts_columns）。
    store_terms = _parse_list(stores)
    if store_terms:
        if len(store_terms) > MAX_STORE_TERMS:
            raise HTTPException(
                400,
                f"门店名称关键词最多 {MAX_STORE_TERMS} 个，当前 {len(store_terms)} 个",
            )
        term_conds = sql.SQL(" OR ").join(
            sql.SQL("{} ILIKE %s").format(sql.Identifier(store_col))
            for _ in store_terms
        )
        conditions.append(sql.SQL("({})").format(term_conds))
        params.extend([f"%{_escape_like(t)}%" for t in store_terms])

    product_list = _parse_list(products)
    if product_list:
        conditions.append(
            sql.SQL("{} IN ({})").format(
                sql.Identifier(product_col),
                sql.SQL(", ").join(sql.Placeholder() * len(product_list)),
            )
        )
        params.extend(product_list)

    where = sql.SQL(" WHERE ") + sql.SQL(" AND ").join(conditions)

    # sp         : 门店 × 品类 盒数聚合
    # st         : 门店总盒数 + 品类数，窗口函数带出全城总量与门店总数
    # top_stores : 按总盒数降序取前 N 家并生成排名
    # sp_rank    : 对前 N 家的每个品类算正序 / 倒序行号
    # 最终只取 rn_desc <= 3 或 rn_asc = 1 的行（Top3 + 末位）
    stmt = sql.SQL(
        """
        WITH sp AS (
            SELECT
                btrim({store})          AS store_name,
                {prod}                  AS product_code,
                MIN({prod_name})        AS product_name,
                SUM(COALESCE({qty}, 0)) AS qty
            FROM {table}{where}
            GROUP BY btrim({store}), {prod}
        ),
        st AS (
            SELECT
                store_name,
                SUM(qty)              AS store_qty,
                COUNT(*)              AS cat_cnt,
                SUM(SUM(qty)) OVER () AS total_qty,
                COUNT(*)      OVER () AS store_count
            FROM sp
            GROUP BY store_name
        ),
        top_stores AS (
            SELECT
                store_name, store_qty, cat_cnt, total_qty, store_count,
                ROW_NUMBER() OVER (ORDER BY store_qty DESC, store_name) AS rank
            FROM st
            ORDER BY store_qty DESC, store_name
            LIMIT %s
        ),
        sp_rank AS (
            SELECT
                sp.store_name, sp.product_code, sp.product_name, sp.qty,
                t.rank, t.store_qty, t.cat_cnt, t.total_qty, t.store_count,
                ROW_NUMBER() OVER (
                    PARTITION BY sp.store_name ORDER BY sp.qty DESC, sp.product_code
                ) AS rn_desc,
                ROW_NUMBER() OVER (
                    PARTITION BY sp.store_name ORDER BY sp.qty ASC, sp.product_code
                ) AS rn_asc
            FROM sp
            JOIN top_stores t ON t.store_name = sp.store_name
        )
        SELECT *
        FROM sp_rank
        WHERE rn_desc <= {top_cat} OR rn_asc = 1
        ORDER BY rank, rn_desc
        """
    ).format(
        store=sql.Identifier(store_col),
        prod=sql.Identifier(product_col),
        prod_name=sql.Identifier(product_name_col),
        qty=sql.Identifier(qty_col),
        table=sql.Identifier(table),
        where=where,
        top_cat=sql.Literal(TOP_CATEGORY_COUNT),
    )

    try:
        with get_conn(db_key) as conn:
            with conn.cursor() as cur:
                cur.execute(stmt, params + [top_n])
                colnames = [d[0] for d in cur.description]
                raw_rows = [dict(zip(colnames, r)) for r in cur.fetchall()]
    except Exception as e:
        logger.error("门店能力分析 %s 失败: %s", db_key, e)
        raise HTTPException(500, f"统计失败: {e}")

    # ---------- 组装结果 ----------
    # 注意：PG 的 SUM/COUNT 经 psycopg2 返回 Decimal / int，且这些值嵌套在 stores 列表的子 dict
    # 中（dict -> list -> dict -> list -> dict）。项目 utils/serialization.jsonable 只递归单值、
    # 不展开嵌套结构；若直接 return，FastAPI 会因 Decimal 不可序列化而 500。故在此统一转
    # float/int（与 jsonable 口径一致），share 因分子分母已是 float 自然也是 float。
    stores: Dict[str, dict] = {}
    total_qty = 0
    store_count = 0

    for r in raw_rows:
        name = r["store_name"]
        store_qty = float(r["store_qty"] or 0)
        qty = float(r["qty"] or 0)
        total_qty = float(r["total_qty"] or 0)
        store_count = int(r["store_count"] or 0)

        s = stores.get(name)
        if s is None:
            s = {
                "rank": int(r["rank"]),
                "store_name": name,
                "qty": store_qty,
                "cat_cnt": int(r["cat_cnt"] or 0),
                "top_categories": [],
                "last_category": None,
            }
            stores[name] = s

        # map_names 开启时：映射表中文名优先，无映射回落到库内商品名称
        code_str = str(r["product_code"]) if r["product_code"] is not None else ""
        display_name = pmap.get(code_str) or r["product_name"] if pmap else r["product_name"]
        item = {
            "product_code": r["product_code"],
            "product_name": display_name,
            "qty": qty,
            "share": (qty / store_qty) if store_qty else None,
        }
        # ORDER BY rank, rn_desc 保证按 1、2、3 顺序追加
        if r["rn_desc"] <= TOP_CATEGORY_COUNT:
            s["top_categories"].append(item)
        if r["rn_asc"] == 1:
            s["last_category"] = item

    store_list = sorted(stores.values(), key=lambda x: x["rank"])

    logger.info(
        "门店能力分析 %s date=[%s,%s] stores=%s products=%s top_n=%s -> 返回 %d 家 / 共 %d 家, 原始行 %d",
        db_key, date_from or "-", date_to or "-", stores or "-", products or "-",
        top_n, len(store_list), store_count, len(raw_rows),
    )

    return {
        "db_key": db_key,
        "city": FIXED_CITY,
        "date_from": date_from,
        "date_to": date_to,
        "summary": {
            "store_count": store_count,   # 广州范围内动销门店总数
            "total_qty": total_qty,       # 广州范围内实销总盒数
            "returned": len(store_list),  # 实际返回门店数（<= top_n）
            "top_n": top_n,
        },
        "stores": store_list,
    }


# ---------- Excel 导出 ----------

def _store_to_row(s: dict) -> list:
    """把单个门店拍平成一行（品类列固定 3 组 Top + 1 组末位）。"""
    line: list = [s.get("rank"), s.get("store_name"), s.get("qty"), s.get("cat_cnt")]
    tops = list(s.get("top_categories") or [])
    for i in range(TOP_CATEGORY_COUNT):
        item = tops[i] if i < len(tops) else None
        line.extend([
            item.get("product_name") if item else "",
            item.get("qty") if item else "",
            item.get("share") if item else "",
        ])
    last = s.get("last_category")
    line.extend([
        last.get("product_name") if last else "",
        last.get("qty") if last else "",
        last.get("share") if last else "",
    ])
    return line


def _store_ability_to_xlsx(result: dict) -> io.BytesIO:
    """把门店能力分析结果写入 xlsx（单个 sheet）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_name("门店能力分析")

    header = ["排名", "门店名称", "实销总数", "品类数"]
    for i in range(1, TOP_CATEGORY_COUNT + 1):
        header.extend([f"最佳品类{i}", f"品类{i}盒数", f"品类{i}占比"])
    header.extend(["末位品类", "末位品类盒数", "末位品类占比"])

    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    ws.freeze_panes = "A2"

    for s in result.get("stores", []):
        ws.append(_store_to_row(s))
        r = ws.max_row
        for c in range(1, len(header) + 1):
            ws.cell(row=r, column=c).font = CELL_FONT

    # 列宽：前两列与品类名列宽一些，数值列适中
    widths = [6, 30, 12, 8]
    for _ in range(TOP_CATEGORY_COUNT + 1):
        widths.extend([26, 12, 10])
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/stats/store_ability")
def store_ability(
    db_key: str = validate_db_key,
    date_from: Optional[str] = Query(None, description="日期起 YYYY-MM-DD（必填；可与 date_to 相同表示单日）"),
    date_to: Optional[str] = Query(None, description="日期止 YYYY-MM-DD（与 date_from 至少填其一）"),
    products: Optional[str] = Query(None, description="品类(商品编码)，英文逗号分隔；留空=全部品类"),
    stores: Optional[str] = Query(None, description="门店名称关键词，英文逗号分隔，模糊匹配；留空=全部门店"),
    top_n: int = Query(DEFAULT_TOP_N, ge=1, le=MAX_TOP_N, description="返回前 N 家门店"),
    map_names: bool = Query(False, description="品类显示映射表中文名（无映射回落商品名称）"),
):
    """门店能力分析（大参林·广州）：总产出排行 + 品类 Top3 与末位。"""
    return _run_store_ability(
        db_key, date_from, date_to, products, top_n, stores=stores, map_names=map_names,
    )


@router.get("/stats/store_ability/export")
def export_store_ability(
    db_key: str = validate_db_key,
    date_from: Optional[str] = Query(None, description="日期起 YYYY-MM-DD（必填）"),
    date_to: Optional[str] = Query(None, description="日期止 YYYY-MM-DD（与 date_from 至少填其一）"),
    products: Optional[str] = Query(None, description="品类(商品编码)，英文逗号分隔；留空=全部品类"),
    stores: Optional[str] = Query(None, description="门店名称关键词，英文逗号分隔，模糊匹配；留空=全部门店"),
    top_n: int = Query(DEFAULT_TOP_N, ge=1, le=MAX_TOP_N),
    map_names: bool = Query(False, description="品类显示映射表中文名（与页面口径一致）"),
):
    """导出门店能力分析结果为 xlsx。"""
    result = _run_store_ability(
        db_key, date_from, date_to, products, top_n, stores=stores, map_names=map_names,
    )
    buf = _store_ability_to_xlsx(result)

    filename = "store_ability_guangzhou.xlsx"
    encoded = quote(filename)
    logger.info(
        "门店能力分析导出 %s date=[%s,%s] stores=%s -> %d 家",
        db_key, date_from or "-", date_to or "-", stores or "-",
        len(result.get("stores", [])),
    )
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )


@router.get("/stats/store_ability/latest_range")
def store_ability_latest_range(db_key: str = validate_db_key):
    """返回广州范围内「数据最新月份」的日期区间，供前端默认填充时间范围。

    销售数据通常滞后于当前自然月（例如现在是 9 月但数据只到 8 月），因此默认区间
    应以库里 MAX(日期) 所在月份为准，而不是当前系统月份，否则默认查出来是空的。
    返回该月的完整首末日；无数据时三个字段均为 null，前端按当月兜底。
    """
    fm = get_field_map(db_key)
    if not fm:
        raise HTTPException(400, f"该统计功能暂不支持数据库: {db_key}")
    if not supports_store_ability(db_key):
        raise HTTPException(400, f"门店能力分析暂不支持数据库: {db_key}（本期仅大参林开放）")

    date_col = fm["date_col"]
    city_col = fm["city_col"]
    table = get_cfg(db_key)["table"]

    # 城市 IN 可命中 (城市, 日期) 索引，MAX 只需定位到索引末尾，代价极低
    stmt = sql.SQL(
        "SELECT MAX({date}) FROM {table} WHERE {city} IN ({ph})"
    ).format(
        date=sql.Identifier(date_col),
        table=sql.Identifier(table),
        city=sql.Identifier(city_col),
        ph=sql.SQL(", ").join(sql.Placeholder() * len(FIXED_CITY_VARIANTS)),
    )

    try:
        with get_conn(db_key) as conn:
            with conn.cursor() as cur:
                cur.execute(stmt, list(FIXED_CITY_VARIANTS))
                row = cur.fetchone()
    except Exception as e:
        logger.error("门店能力分析 取最新月份失败 %s: %s", db_key, e)
        raise HTTPException(500, f"获取最新月份失败: {e}")

    max_date = row[0] if row else None
    if max_date is None:
        return {"date_from": None, "date_to": None, "max_date": None}

    first = max_date.replace(day=1)
    last_day = calendar.monthrange(max_date.year, max_date.month)[1]
    last = max_date.replace(day=last_day)
    return {
        "date_from": first.isoformat(),
        "date_to": last.isoformat(),
        "max_date": max_date.isoformat(),
    }


@router.get("/stats/store_ability/store_trend")
def store_ability_trend(
    db_key: str = validate_db_key,
    store_name: str = Query(..., min_length=1, description="门店名称（表格中的完整名称，精确匹配）"),
    date_from: str = Query(..., description="日期起 YYYY-MM-DD（必填）"),
    date_to: str = Query(..., description="日期止 YYYY-MM-DD（必填）"),
    products: Optional[str] = Query(None, description="品类(商品编码)，英文逗号分隔；留空=全部"),
):
    """单个门店在日期区间内的「按天产出」数据。

    - 按天聚合实销盒数；区间内没有销量的日期补 0，保证柱状图横轴连续。
    - 不再区分「完整月 / 非完整月」：按天呈现时每一天本身就是完整的，
      首尾日期不会被截断，因此不存在部分月份的口径问题。
    - products 与主表共用同一筛选口径，保证「表格里的实销总数 = 图表各日合计」。
    """
    fm = get_field_map(db_key)
    if not fm:
        raise HTTPException(400, f"该统计功能暂不支持数据库: {db_key}")
    if not supports_store_ability(db_key):
        raise HTTPException(400, f"门店能力分析暂不支持数据库: {db_key}（本期仅大参林开放）")

    try:
        d_from = datetime.strptime(date_from, "%Y-%m-%d").date()
        d_to = datetime.strptime(date_to, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(400, "日期格式应为 YYYY-MM-DD")
    if d_from > d_to:
        raise HTTPException(400, "date_from 不能晚于 date_to")

    # 按天聚合的柱子上限：防止误选超长区间（如十年）导致柱子密到无法阅读
    day_count = (d_to - d_from).days + 1
    if day_count > MAX_TREND_DAYS:
        raise HTTPException(
            400,
            f"日期区间过长（{day_count} 天），按天展示上限为 {MAX_TREND_DAYS} 天，请缩小范围",
        )

    store_col = fm["store_name_col"]
    product_col = fm["product_col"]
    qty_col = fm["qty_col"]
    city_col = fm["city_col"]
    date_col = fm["date_col"]
    table = get_cfg(db_key)["table"]

    # 条件顺序必须与 params 追加顺序一致
    conditions: List[sql.Composable] = [
        sql.SQL("{} IN ({})").format(
            sql.Identifier(city_col),
            sql.SQL(", ").join(sql.Placeholder() * len(FIXED_CITY_VARIANTS)),
        ),
        # 表格里的 store_name 已是 btrim 后的值，这里精确匹配即可（非模糊）
        sql.SQL("btrim({}) = %s").format(sql.Identifier(store_col)),
        sql.SQL("{} >= %s").format(sql.Identifier(date_col)),
        sql.SQL("{} <= %s").format(sql.Identifier(date_col)),
    ]
    params: list = [*FIXED_CITY_VARIANTS, store_name.strip(), date_from, date_to]

    product_list = _parse_list(products)
    if product_list:
        conditions.append(
            sql.SQL("{} IN ({})").format(
                sql.Identifier(product_col),
                sql.SQL(", ").join(sql.Placeholder() * len(product_list)),
            )
        )
        params.extend(product_list)

    where = sql.SQL(" WHERE ") + sql.SQL(" AND ").join(conditions)
    stmt = sql.SQL(
        "SELECT to_char({date}, 'YYYY-MM-DD') AS day, "
        "SUM(COALESCE({qty}, 0)) AS qty "
        "FROM {table}{where} "
        "GROUP BY 1 ORDER BY 1"
    ).format(
        date=sql.Identifier(date_col),
        qty=sql.Identifier(qty_col),
        table=sql.Identifier(table),
        where=where,
    )

    try:
        with get_conn(db_key) as conn:
            with conn.cursor() as cur:
                cur.execute(stmt, params)
                colnames = [d[0] for d in cur.description]
                raw_rows = [dict(zip(colnames, r)) for r in cur.fetchall()]
    except Exception as e:
        logger.error("门店月度趋势 %s/%s 失败: %s", db_key, store_name, e)
        raise HTTPException(500, f"统计失败: {e}")

    # Decimal -> float；再用日期序列补齐区间内没有销量的日期（qty=0）
    qty_by_day = {r["day"]: float(r["qty"] or 0) for r in raw_rows}
    days = [
        {"date": d, "qty": qty_by_day.get(d, 0)}
        for d in _day_sequence(d_from, d_to)
    ]

    logger.info(
        "门店每日趋势 %s store=%s date=[%s,%s] products=%s -> %d 天",
        db_key, store_name, date_from, date_to, products or "-", len(days),
    )

    return {
        "db_key": db_key,
        "store_name": store_name.strip(),
        "date_from": date_from,
        "date_to": date_to,
        "days": days,
    }
