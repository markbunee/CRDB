# -*- coding: utf-8 -*-
"""门店数统计接口：/api/{db_key}/stats/store_count(+ /export)

按维度统计实销门店数（DISTINCT 门店编码 非重复计数）。

统一为「拆分维度(多表) + 行维度 + 列维度」模型，支持三种维度：
  - dimension=city    地域维度：拆分=地域(可合并)，行=品类(商品编码)，列=月份(可合并)
  - dimension=product 品类维度：拆分=品类(可合并)，行=地域，列=月份(可合并)
  - dimension=time    时间维度：拆分=月份(可合并)，行=地域，列=品类

地域级别 region_level 可在城市(city)/省份(province)间二选一切换（城市与省份互斥）。

导出：GET /stats/store_count/export 返回 xlsx（每个 table 一个 sheet），英文文件名。
"""
import io
import re
from typing import Any, Dict, List, Optional
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from psycopg2 import sql

from ..dependencies import validate_db_key
from ..models.schema_def import get_cfg, get_field_map, get_region_levels, get_region_label
from ..services.product_map import get_product_map
from ..database import get_conn
from ..logger import get_logger

router = APIRouter(prefix="/api/{db_key}", tags=["stats-store"])
logger = get_logger("app.routers.stats_store")

DIMENSIONS = ("city", "product", "time")

# ---------- Excel 样式 ----------
HEADER_FILL = PatternFill(start_color="3B6BD6", end_color="3B6BD6", fill_type="solid")
HEADER_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Microsoft YaHei", size=10)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)


def _parse_list(val: Optional[str]) -> List[str]:
    if not val:
        return []
    return [v.strip() for v in val.split(",") if v.strip() != ""]


def _fmt_month(val: Any) -> str:
    if val is None:
        return "未知"
    if hasattr(val, "strftime"):
        return val.strftime("%Y-%m")
    s = str(val)
    return s[:7] if len(s) >= 7 else s


def _val_to_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val)


def _fmt_val(fm: Dict[str, str], col_name: Optional[str], val: Any,
             pmap: Optional[Dict[str, str]] = None) -> Optional[str]:
    if col_name is None or val is None:
        return None
    if col_name == fm["month_col"]:
        return _fmt_month(val)
    s = _val_to_str(val)
    # 品类编码 → 品类中文名（仅 map_names 开启时 pmap 非空；无映射的编码保持原样）
    if pmap and col_name == fm["product_col"]:
        return pmap.get(s, s)
    return s


def _range_label(date_from: Optional[str], date_to: Optional[str]) -> str:
    if date_from and date_to:
        return f"{date_from} ~ {date_to}"
    if date_from:
        return f"{date_from} ~ 至今"
    if date_to:
        return f"截至 {date_to}"
    return "全部日期"


def _safe_sort(vals: list) -> list:
    def key(v):
        if v is None:
            return (1, 0, "")
        return (0, 0, v)
    try:
        return sorted(vals, key=key)
    except TypeError:
        return sorted(vals, key=lambda v: (v is None, str(v)))


def _build_where(
    fm: Dict[str, str],
    date_from: Optional[str],
    date_to: Optional[str],
    product_list: List[str],
    region_col: str,
    region_list: List[str],
):
    conditions: List[sql.Composable] = []
    params: list = []

    if date_from:
        conditions.append(sql.SQL("{} >= %s").format(sql.Identifier(fm["date_col"])))
        params.append(date_from)
    if date_to:
        conditions.append(sql.SQL("{} <= %s").format(sql.Identifier(fm["date_col"])))
        params.append(date_to)

    if product_list:
        placeholders = sql.SQL(", ").join(sql.Placeholder() * len(product_list))
        conditions.append(
            sql.SQL("{} IN ({})").format(sql.Identifier(fm["product_col"]), placeholders)
        )
        params.extend(product_list)

    if region_list:
        placeholders = sql.SQL(", ").join(sql.Placeholder() * len(region_list))
        conditions.append(
            sql.SQL("{} IN ({})").format(sql.Identifier(region_col), placeholders)
        )
        params.extend(region_list)

    where = sql.SQL("")
    if conditions:
        where = sql.SQL(" WHERE ") + sql.SQL(" AND ").join(conditions)
    return where, params


def _make_title(
    dimension: str,
    split_label: Optional[str],
    merge_split: bool,
    date_from: Optional[str],
    date_to: Optional[str],
    split_name: str,
    region_name: str,
) -> str:
    if merge_split or split_label is None:
        if dimension == "city":
            return f"全部{region_name}（合并）"
        if dimension == "product":
            return "全部品类（合并）"
        return f"全部月份（合并）· {_range_label(date_from, date_to)}"
    return f"{split_name}: {split_label}"


def _run_store_count(
    db_key: str,
    dimension: str,
    region_level: str,
    date_from: Optional[str],
    date_to: Optional[str],
    merge_months: bool,
    cities: Optional[str],
    provinces: Optional[str],
    products: Optional[str],
    merge_cities: bool,
    merge_products: bool,
    map_names: bool = False,
) -> dict:
    """门店数统计核心逻辑：查询 + 透视，返回结果 dict（JSON 端点与导出共用）。

    map_names=True 时把品类编码翻译成映射表中文名（无映射保持编码），
    JSON 与导出共用本函数，因此两者的口径天然一致。
    """
    fm = get_field_map(db_key)
    if not fm:
        raise HTTPException(400, f"该统计功能暂不支持数据库: {db_key}")
    if dimension not in DIMENSIONS:
        raise HTTPException(400, f"不支持的维度: {dimension}，可选: {DIMENSIONS}")

    allowed_regions = get_region_levels(db_key)
    if region_level not in allowed_regions:
        raise HTTPException(400, f"不支持的地域级别: {region_level}，可选: {allowed_regions}")

    region_label = get_region_label(db_key)
    cfg = get_cfg(db_key)
    table = cfg["table"]

    if region_level == "province" and "province_col" in fm:
        region_col = fm["province_col"]
        region_name = "省份"
        region_list = _parse_list(provinces)
    else:
        region_col = fm["city_col"]
        region_name = region_label
        region_list = _parse_list(cities)

    product_list = _parse_list(products)

    if dimension == "city":
        split_col = region_col
        row_col = fm["product_col"]
        col_col = fm["month_col"]
        merge_split = merge_cities
        split_name, row_name, col_name = region_name, "品类", "月份"
    elif dimension == "product":
        split_col = fm["product_col"]
        row_col = region_col
        col_col = fm["month_col"]
        merge_split = merge_products
        split_name, row_name, col_name = "品类", region_name, "月份"
    else:
        split_col = fm["month_col"]
        row_col = region_col
        col_col = fm["product_col"]
        merge_split = merge_months
        split_name, row_name, col_name = "月份", region_name, "品类"

    month_col = fm["month_col"]

    # 根据当前维度，确定 split/row/col 分别由哪个 merge 参数控制
    if dimension == "city":
        merge_for: Dict[str, bool] = {
            split_col: merge_cities,
            row_col: merge_products,
            col_col: merge_months,
        }
    elif dimension == "product":
        merge_for = {
            split_col: merge_products,
            row_col: merge_cities,
            col_col: merge_months,
        }
    else:
        merge_for = {
            split_col: merge_months,
            row_col: merge_cities,
            col_col: merge_products,
        }

    group_cols: List[sql.Identifier] = []
    for col, merge_flag in merge_for.items():
        if col and not merge_flag:
            group_cols.append(sql.Identifier(col))

    where, params = _build_where(fm, date_from, date_to, product_list, region_col, region_list)

    fields = list(group_cols) + [
        sql.SQL("COUNT(DISTINCT {}) AS store_count").format(
            sql.Identifier(fm["store_col"])
        )
    ]
    if group_cols:
        stmt = sql.SQL(
            "SELECT {fields} FROM {table}{where} GROUP BY {group} ORDER BY {order}"
        ).format(
            fields=sql.SQL(", ").join(fields),
            table=sql.Identifier(table),
            where=where,
            group=sql.SQL(", ").join(group_cols),
            order=sql.SQL(", ").join(group_cols),
        )
    else:
        stmt = sql.SQL(
            "SELECT {fields} FROM {table}{where}"
        ).format(
            fields=sql.SQL(", ").join(fields),
            table=sql.Identifier(table),
            where=where,
        )

    try:
        with get_conn(db_key) as conn:
            with conn.cursor() as cur:
                cur.execute(stmt, params)
                colnames = [d[0] for d in cur.description]
                raw_rows = [dict(zip(colnames, r)) for r in cur.fetchall()]
    except Exception as e:
        logger.error("门店数统计 %s 失败: %s", db_key, e)
        raise HTTPException(500, f"统计失败: {e}")

    # 品类编码映射：仅在开关打开时读盘（文件很小，读一次足够）
    pmap: Optional[Dict[str, str]] = get_product_map(db_key) if map_names else None

    tables = _pivot(
        raw_rows, fm, dimension, merge_for,
        date_from, date_to,
        split_col, row_col, col_col,
        split_name, row_name, col_name, region_name,
        pmap,
    )

    logger.info(
        "门店数统计 %s dim=%s region=%s date=[%s,%s] merge_months=%s regions=%s products=%s "
        "merge_cities=%s merge_products=%s -> %d 表, %d 原始行",
        db_key, dimension, region_level, date_from or "-", date_to or "-",
        merge_months, cities or "-" if region_level == "city" else provinces or "-",
        products or "-", merge_cities, merge_products, len(tables), len(raw_rows),
    )

    return {
        "tables": tables,
        "dimension": dimension,
        "region_level": region_level,
        "merge_months": merge_months,
        "merge_cities": merge_cities,
        "merge_products": merge_products,
        "total_raw_rows": len(raw_rows),
    }


def _pivot(
    raw_rows: List[dict],
    fm: Dict[str, str],
    dimension: str,
    merge_for: Dict[str, bool],
    date_from: Optional[str],
    date_to: Optional[str],
    split_col: Optional[str],
    row_col: Optional[str],
    col_col: Optional[str],
    split_name: str,
    row_name: str,
    col_name: str,
    region_name: str,
    pmap: Optional[Dict[str, str]] = None,
) -> List[dict]:
    month_col = fm["month_col"]
    merge_split = merge_for.get(split_col, False) if split_col else False
    merge_row = merge_for.get(row_col, False) if row_col else False
    merge_col = merge_for.get(col_col, False) if col_col else False
    col_is_month_merged = (col_col == month_col and merge_col)

    split_raw: list = []
    split_seen: set = set()
    row_raw: list = []
    row_seen: set = set()
    col_raw: list = []
    col_seen: set = set()

    for r in raw_rows:
        if split_col and not merge_split:
            v = r.get(split_col)
            k = _val_to_str(v)
            if k and k not in split_seen:
                split_seen.add(k)
                split_raw.append(v)
        if row_col and not merge_row:
            v = r.get(row_col)
            k = _val_to_str(v)
            if k and k not in row_seen:
                row_seen.add(k)
                row_raw.append(v)
        if col_col and not merge_col:
            v = r.get(col_col)
            k = _val_to_str(v)
            if k and k not in col_seen:
                col_seen.add(k)
                col_raw.append(v)

    if merge_split:
        split_labels: List[Optional[str]] = [None]
    else:
        split_labels = [_fmt_val(fm, split_col, v, pmap) for v in _safe_sort(split_raw)]

    if merge_row:
        row_keys = [f"全部{row_name}（合并）"]
    else:
        row_keys = [_fmt_val(fm, row_col, v, pmap) for v in _safe_sort(row_raw)] if row_col else []

    if col_col:
        if col_is_month_merged:
            col_labels = [_range_label(date_from, date_to)]
        elif merge_col:
            col_labels = [f"全部{col_name}（合并）"]
        else:
            col_labels = [_fmt_val(fm, col_col, v, pmap) for v in _safe_sort(col_raw)]
    else:
        col_labels = []

    index: Dict[tuple, int] = {}
    for r in raw_rows:
        count = r.get("store_count", 0)
        if merge_split:
            sl: Optional[str] = None
        else:
            sl = _fmt_val(fm, split_col, r.get(split_col), pmap)
        if merge_row:
            rk = f"全部{row_name}（合并）"
        else:
            rk = _fmt_val(fm, row_col, r.get(row_col), pmap) if row_col else ""
        if col_col:
            if col_is_month_merged:
                cl = _range_label(date_from, date_to)
            elif merge_col:
                cl = f"全部{col_name}（合并）"
            else:
                cl = _fmt_val(fm, col_col, r.get(col_col), pmap)
        else:
            cl = ""
        index[(sl, rk, cl)] = count

    tables: List[dict] = []
    for sl in split_labels:
        rows = []
        for rk in row_keys:
            cells = {cl: index.get((sl, rk, cl), 0) for cl in col_labels}
            rows.append({"row_key": rk, "cells": cells})
        tables.append(
            {
                "title": _make_title(dimension, sl, merge_split, date_from, date_to, split_name, region_name),
                "split_value": sl,
                "row_header": row_name,
                "col_header": col_name,
                "row_keys": row_keys,
                "col_keys": col_labels,
                "rows": rows,
            }
        )

    return tables


# ---------- Excel 导出 ----------

def _safe_sheet_name(title: str, idx: int) -> str:
    """生成合法 Excel sheet 名（去非法字符、截断 31 字符）。"""
    name = re.sub(r'[:\\/\?\*\[\]]', "_", str(title)).strip()
    if not name:
        name = f"sheet{idx + 1}"
    if len(name) > 31:
        name = name[:31]
    return name


def _tables_to_xlsx(tables: List[dict]) -> io.BytesIO:
    """把多张透视表写入一个 xlsx（每个 table 一个 sheet）。"""
    wb = Workbook()
    wb.remove(wb.active)  # 删除默认 sheet

    for i, t in enumerate(tables):
        ws = wb.create_sheet(title=_safe_sheet_name(t.get("title") or f"sheet{i+1}", i))
        col_keys = list(t.get("col_keys", []))
        row_keys = list(t.get("row_keys", []))
        rows = t.get("rows", [])

        # 表头：行表头 + 各列
        header = [t.get("row_header", "")] + col_keys
        ws.append(header)
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGN
            cell.border = THIN_BORDER
        ws.freeze_panes = "A2"

        # 数据行
        for row in rows:
            line = [row.get("row_key", "")]
            cells = row.get("cells", {})
            for ck in col_keys:
                v = cells.get(ck, 0)
                line.append(v if v is not None else 0)
            ws.append(line)
            # 普通单元格字体
            r = ws.max_row
            for c in range(1, len(line) + 1):
                ws.cell(row=r, column=c).font = CELL_FONT

        # 列宽
        ws.column_dimensions["A"].width = 18
        for j in range(len(col_keys)):
            ws.column_dimensions[get_column_letter(j + 2)].width = 16

    if not wb.worksheets:
        ws = wb.create_sheet(title="empty")
        ws.append(["无数据"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/stats/store_count")
def store_count(
    db_key: str = validate_db_key,
    dimension: str = Query("city", description="统计维度: city | product | time"),
    region_level: str = Query("city", description="地域级别: city | province"),
    date_from: Optional[str] = Query(None, description="日期起 YYYY-MM-DD"),
    date_to: Optional[str] = Query(None, description="日期止 YYYY-MM-DD"),
    merge_months: bool = Query(False, description="合并月范围"),
    cities: Optional[str] = Query(None, description="城市，英文逗号分隔（region_level=city 时生效）"),
    provinces: Optional[str] = Query(None, description="省份，英文逗号分隔（region_level=province 时生效）"),
    products: Optional[str] = Query(None, description="品类(商品编码)，英文逗号分隔"),
    merge_cities: bool = Query(False, description="合并地域维度为单表"),
    merge_products: bool = Query(False, description="合并品类维度为单表"),
    map_names: bool = Query(False, description="品类编码映射为中文名（无映射保持编码）"),
):
    """实销门店数统计 — 通用三维度 + 地域级别切换。"""
    return _run_store_count(
        db_key, dimension, region_level, date_from, date_to, merge_months,
        cities, provinces, products, merge_cities, merge_products, map_names,
    )


@router.get("/stats/store_count/export")
def export_store_count(
    db_key: str = validate_db_key,
    dimension: str = Query("city", description="统计维度: city | product | time"),
    region_level: str = Query("city", description="地域级别: city | province"),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    merge_months: bool = Query(False),
    cities: Optional[str] = Query(None),
    provinces: Optional[str] = Query(None),
    products: Optional[str] = Query(None),
    merge_cities: bool = Query(False),
    merge_products: bool = Query(False),
    map_names: bool = Query(False, description="品类编码映射为中文名（与页面口径一致）"),
):
    """导出门店数统计结果为 xlsx（每个 table 一个 sheet，英文文件名）。"""
    result = _run_store_count(
        db_key, dimension, region_level, date_from, date_to, merge_months,
        cities, provinces, products, merge_cities, merge_products, map_names,
    )
    tables = result["tables"]
    buf = _tables_to_xlsx(tables)

    filename = f"store_count_{dimension}_{region_level}.xlsx"
    encoded = quote(filename)
    logger.info(
        "门店数导出 %s dim=%s region=%s -> %d 表",
        db_key, dimension, region_level, len(tables),
    )
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )
